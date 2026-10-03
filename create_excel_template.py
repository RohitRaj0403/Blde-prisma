import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_template():
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    font_family = "Segoe UI"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    title_font = Font(name=font_family, size=14, bold=True, color="1E3A8A")
    normal_font = Font(name=font_family, size=10)
    bold_font = Font(name=font_family, size=10, bold=True)
    italic_font = Font(name=font_family, size=10, italic=True)

    blue_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    green_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    gray_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    header_border = Border(
        left=Side(style='thin', color='FFFFFF'),
        right=Side(style='thin', color='FFFFFF'),
        top=Side(style='medium', color='1E3A8A'),
        bottom=Side(style='medium', color='1E3A8A')
    )

    # ──── 1. FACULTY PROFILES SHEET ───────────────────────────────────────────
    ws_faculty = wb.create_sheet(title="Faculty Profiles")
    ws_faculty.views.sheetView[0].showGridLines = True
    
    faculty_headers = [
        "fullName", "email", "department", "orcid", "scopusAuthorId", 
        "googleScholarId", "vidwanId", "researcherId", "bio", 
        "education", "experience", "expertise", "photoUrl"
    ]
    
    # Header styling
    for col_idx, h in enumerate(faculty_headers, 1):
        cell = ws_faculty.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = green_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    # Sample Data
    faculty_samples = [
        [
            "Dr. Anil Kumar S",
            "faculty@blde.ac.in",
            "Medicine",
            "0000-0001-1234-5678",
            "57195432101",
            "fo42anktZ9vv",
            "394851",
            "H-8274-2018",
            "Professor & Head of Medicine at BLDE Medical College. Specializes in cardiovascular diseases and clinical research methodologies.",
            "• Ph.D. in Medicine, Bangalore Medical College, 2012\n• MD in General Medicine, KIMS Hubli, 2008\n• MBBS, Mysore Medical College, 2004",
            "• Professor & Head, Department of Medicine, BLDE DU (2020 - Present)\n• Associate Professor, Department of Medicine, BLDE DU (2015 - 2020)",
            "Cardiology, Diabetes Mellitus, Clinical Trials",
            "https://images.unsplash.com/photo-1622253692010-333f2da6031d"
        ],
        [
            "Dr. Priya Sharma",
            "priya.sharma@blde.ac.in",
            "Biochemistry",
            "0000-0002-9876-5432",
            "57234567890",
            "xPq8901tZ2vv",
            "394852",
            "H-8274-2019",
            "Associate Professor of Biochemistry with research focus on metabolic pathways.",
            "• MD in Biochemistry, JIPMER, 2013\n• MBBS, Kasturba Medical College, 2009",
            "• Associate Professor, BLDE DU (2018 - Present)\n• Assistant Professor, Manipal Academy of Higher Education (2014 - 2018)",
            "Biochemistry, Endocrinology, Biomarkers",
            "https://images.unsplash.com/photo-1594824813573-246434de83fb"
        ]
    ]

    for row_idx, row_data in enumerate(faculty_samples, 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_faculty.cell(row=row_idx, column=col_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if col_idx in [10, 11]:  # Wrap text for multiline text areas (education/experience)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top")

    # ──── 2. PUBLICATIONS SHEET ───────────────────────────────────────────────
    ws_pubs = wb.create_sheet(title="Publications")
    ws_pubs.views.sheetView[0].showGridLines = True
    
    pub_headers = [
        "title", "year", "doi", "type", "journal", "issn", 
        "abstract", "keywords", "citations", "pmid", "wosUid", "scopusEid"
    ]
    
    # Header styling
    for col_idx, h in enumerate(pub_headers, 1):
        cell = ws_pubs.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = blue_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    # Sample Data
    pub_samples = [
        [
            "Prevalence of Type 2 Diabetes and its Association with Metabolic Syndrome in Rural Karnataka",
            2023,
            "10.1234/ijmr.2023.001",
            "Article",
            "Indian Journal of Medical Research",
            "0971-5916",
            "A cross-sectional study examining the prevalence of Type 2 diabetes and metabolic syndrome in rural Karnataka.",
            "Diabetes Mellitus; Metabolic Syndrome; Karnataka",
            12,
            "",
            "",
            "57195432101"
        ],
        [
            "Histopathological Patterns of Cervical Lesions: A Five-Year Retrospective Study",
            2022,
            "10.1234/ijpm.2022.045",
            "Article",
            "Indian Journal of Pathology and Microbiology",
            "0377-4929",
            "Retrospective analysis of cervical biopsy specimens over five years at a tertiary care center.",
            "Cervical Pathology; Histopathology",
            7,
            "34567890",
            "",
            ""
        ]
    ]

    for row_idx, row_data in enumerate(pub_samples, 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_pubs.cell(row=row_idx, column=col_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=(col_idx in [1, 7]))

    # ──── 3. GUIDELINES & FORMATTING SHEET ────────────────────────────────────
    ws_guide = wb.create_sheet(title="Guidelines & Formatting")
    ws_guide.views.sheetView[0].showGridLines = True
    
    ws_guide.cell(row=1, column=1, value="BLDE PRISM - Dataset Import Guidelines").font = title_font
    
    guide_headers = ["Field Name", "Applies To", "Description / Formatting Rules", "Example Value"]
    for col_idx, h in enumerate(guide_headers, 1):
        cell = ws_guide.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = blue_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = header_border

    guide_rows = [
        ["fullName", "Faculty Profiles", "Full name of the faculty member (Required).", "Dr. Anil Kumar S"],
        ["email", "Faculty Profiles", "Institutional email. Used to link profiles with their user login accounts (Optional).", "faculty@blde.ac.in"],
        ["department", "Faculty Profiles", "Must match one of the system departments (Medicine, Surgery, Biochemistry, etc.).", "Medicine"],
        ["orcid / scopusAuthorId / googleScholarId / vidwanId / researcherId", "Faculty Profiles", "Academic identifier numbers. Used for profile link widgets (Optional).", "orcid: 0000-0001-1234-5678\nvidwanId: 394851"],
        ["bio", "Faculty Profiles", "Short biography text (Optional).", "Professor & Head of Medicine..."],
        ["education", "Faculty Profiles", "Multi-line text detailing education. Start each entry with a bullet point (•) on a new line (Alt+Enter in Excel) to automatically generate the timeline visualization.", "• Ph.D., Bangalore College, 2012\n• MD, KIMS Hubli, 2008"],
        ["experience", "Faculty Profiles", "Multi-line text detailing work history. Format with bullet points (•) on separate lines (Alt+Enter) to generate the timeline visualization.", "• Professor, BLDE DU (2020 - Pres.)\n• Associate Prof., BLDE (2015-2020)"],
        ["expertise", "Faculty Profiles", "Comma-separated list of research area tags.", "Cardiology, Diabetes, Clinical Trials"],
        ["photoUrl", "Faculty Profiles", "Fully qualified web URL to the profile photo (Optional).", "https://images.unsplash.com/photo-1622253692010-333f2da6031d"],
        ["", "", "", ""], # Separator
        ["title", "Publications", "The title of the publication (Required).", "Prevalence of Type 2 Diabetes..."],
        ["year", "Publications", "Publication year as a 4-digit number (Required).", "2023"],
        ["doi", "Publications", "Digital Object Identifier (Recommended for dedup checking).", "10.1234/ijmr.2023.001"],
        ["type", "Publications", "Publication document type. Maps to: Article, Review, Case Report, Book Chapter, Conference Paper, Letter, Editorial, Other.", "Article"],
        ["journal", "Publications", "Journal name where the paper was published.", "Indian Journal of Medical Research"],
        ["issn", "Publications", "International Standard Serial Number of the journal.", "0971-5916"],
        ["keywords", "Publications", "Keywords separated by semicolon (;).", "Diabetes Mellitus; Metabolic Syndrome"],
        ["citations", "Publications", "Total citations integer count.", "12"]
    ]

    for row_idx, row_data in enumerate(guide_rows, 4):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_guide.cell(row=row_idx, column=col_idx, value=val)
            if col_idx == 1 and val:
                cell.font = bold_font
            else:
                cell.font = normal_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # ──── Auto-fit columns for all sheets ─────────────────────────────────────
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            
            # Special case for headers in guidelines
            if sheet.title == "Guidelines & Formatting" and col[0].column == 3:
                sheet.column_dimensions[col_letter].width = 60
                continue

            for cell in col:
                if cell.value:
                    # Handle multi-line strings length check
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        if len(line) > max_len:
                            max_len = len(line)
            
            # Set width with a safety margin
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save Excel Workbook
    output_path = "blde_prism_import_template.xlsx"
    wb.save(output_path)
    print(f"Excel template generated successfully at: {output_path}")

if __name__ == "__main__":
    create_template()
